from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from scraper.models import PlaceRecord
from scraper.storage import atomic_save_workbook

BASE_COLUMNS = [
    "firmenname",
    "strasse",
    "hausnummer",
    "plz",
    "ort",
    "bundesland",
    "telefon",
    "telefon_international",
    "website",
    "website_quelle",
    "email",
    "email_quelle_url",
    "email_typ",
    "email_confidence",
    "google_places_id",
    "google_maps_link",
    "google_primary_type",
    "google_types",
    "business_status",
    "breitengrad",
    "laengengrad",
    "erstfund_keyword",
    "gefunden_durch_suchbegriffe",
    "anzahl_suchbegriffe",
    "quelle",
    "relevanz_score",
    "relevanz_status",
    "klassifizierung",
    "klassifizierung_score",
    "klassifizierung_grund",
    "scrape_datum",
]
MANUAL_COLUMNS = ["pruefgrund", "pruefhinweise"]
DISCARDED_COLUMNS = ["aussortiert_grund", "aussortiert_details", "duplikat_von"]


def _columns(search_terms: list[str], extra: list[str] | None = None) -> list[str]:
    keyword_columns = ["kw_" + term.lower().replace(" ", "_").replace("-", "_") for term in search_terms]
    return BASE_COLUMNS[:24] + keyword_columns + BASE_COLUMNS[24:] + (extra or [])


def _keyword_overlap_columns(search_terms: list[str]) -> list[str]:
    return ["auch_kw_" + term.lower().replace(" ", "_").replace("-", "_") for term in search_terms]


def _write_workbook(path: Path, sheet_name: str, rows: list[dict[str, Any]], columns: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([row.get(column, "") for column in columns])
    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"
    text_columns = {"plz", "telefon", "telefon_international", "email", "google_places_id"}
    for idx, column in enumerate(columns, start=1):
        letter = get_column_letter(idx)
        max_len = max([len(str(column))] + [len(str(row.get(column, ""))) for row in rows])
        sheet.column_dimensions[letter].width = min(max(max_len + 2, 12), 48)
        if column in text_columns:
            for cell in sheet[letter][1:]:
                cell.number_format = "@"
        if column.startswith("kw_"):
            for cell in sheet[letter][1:]:
                cell.value = bool(cell.value)
    atomic_save_workbook(workbook, path)


def export_final_workbooks(
    komplett: list[PlaceRecord],
    manuell: list[PlaceRecord],
    aussortiert: list[PlaceRecord],
    search_terms: list[str],
    final_dir: str | Path = "data/final",
) -> None:
    final_path = Path(final_dir)
    _write_workbook(
        final_path / "alles_komplett.xlsx",
        "alles_komplett",
        [record.to_row(search_terms) for record in komplett],
        _columns(search_terms),
    )
    _write_workbook(
        final_path / "manuelle_pruefung.xlsx",
        "manuelle_pruefung",
        [record.to_row(search_terms) for record in manuell],
        _columns(search_terms, MANUAL_COLUMNS),
    )
    _write_workbook(
        final_path / "aussortiert.xlsx",
        "aussortiert",
        [record.to_row(search_terms) for record in aussortiert],
        _columns(search_terms, DISCARDED_COLUMNS),
    )


def export_keyword_statistics(rows: list[dict[str, Any]], path: str | Path, search_terms: list[str] | None = None) -> None:
    columns = [
        "suchbegriff",
        "rohtreffer",
        "eindeutige_place_ids",
        "exklusive_place_ids",
        "auch_ueber_andere_keywords",
        "auch_ueber_andere_keywords_details",
        "keyword_kombinationen_details",
        "anteil_exklusiv_prozent",
    ]
    if search_terms:
        columns.extend(_keyword_overlap_columns(search_terms))
    _write_workbook(Path(path), "keyword_statistik", rows, columns)


def verify_workbook(path: str | Path, required_columns: list[str]) -> bool:
    workbook = load_workbook(path)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    return all(column in headers for column in required_columns)
