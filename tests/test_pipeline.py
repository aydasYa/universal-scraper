from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from scraper.config import BranchConfig
from scraper.costs import estimate_google_places_cost
from scraper.discovery.google_places import CheckpointStore, GooglePlacesClient, Tile, discover_google_places
from scraper.enrichment.contacts import EmailCandidate, enrich_contacts
from scraper.export.xlsx import verify_workbook
from scraper.matching.deduplicate import deduplicate_records
from scraper.matching.normalize import normalize_domain, normalize_phone, normalize_url
from scraper.matching.relevance import classify_record
from scraper.models import PlaceRecord
from scraper.pipeline import build_keyword_statistics, load_raw_records, process_raw_to_outputs


TERMS = ["Herrenfriseur", "Barbershop", "Damenfriseur", "Friseursalon"]
CONFIG = BranchConfig(slug="friseure", name="Friseure", search_terms=TERMS)


def rec(place_id: str, name: str, keyword: str = "Friseursalon", address: str = "Musterstr 1, 10115 Berlin", website: str = "") -> PlaceRecord:
    return PlaceRecord(
        firmenname=name,
        strasse="musterstr",
        hausnummer="1",
        plz="10115",
        ort="berlin",
        website=website,
        google_places_id=place_id,
        google_primary_type="hair_salon",
        google_types="hair_salon, establishment",
        raw_keywords={keyword},
        erstfund_keyword=keyword,
    )


def raw(place_id: str, name: str, keyword: str, types: list[str] | None = None, website: str = "") -> dict:
    return {
        "keyword": keyword,
        "search_result": {
            "id": place_id,
            "name": f"places/{place_id}",
            "displayName": {"text": name, "languageCode": "de"},
            "formattedAddress": "Musterstr 1, 10115 Berlin",
        },
        "details": {
            "id": place_id,
            "name": f"places/{place_id}",
            "displayName": {"text": name, "languageCode": "de"},
            "formattedAddress": "Musterstr 1, 10115 Berlin",
            "addressComponents": [
                {"longText": "10115", "types": ["postal_code"]},
                {"longText": "Berlin", "types": ["locality"]},
                {"longText": "Berlin", "types": ["administrative_area_level_1"]},
            ],
            "websiteUri": website,
            "types": types or ["hair_salon", "establishment"],
            "primaryType": (types or ["hair_salon"])[0],
            "businessStatus": "OPERATIONAL",
            "location": {"latitude": 52.5, "longitude": 13.4},
        },
    }


class PipelineTests(unittest.TestCase):
    def test_google_places_new_text_search_request_shape(self) -> None:
        class Response:
            status_code = 200

            def raise_for_status(self) -> None:
                return None

            def json(self) -> dict:
                return {"places": [{"id": "pid-1", "displayName": {"text": "Salon"}}]}

        with patch("scraper.discovery.google_places.requests.post", return_value=Response()) as post:
            client = GooglePlacesClient(api_key="test-key")
            places = client.text_search("Friseursalon", Tile(52.4, 13.3, 52.5, 13.4))

        self.assertEqual(places[0]["id"], "pid-1")
        _, kwargs = post.call_args
        self.assertEqual(kwargs["headers"]["X-Goog-Api-Key"], "test-key")
        self.assertIn("places.websiteUri", kwargs["headers"]["X-Goog-FieldMask"])
        self.assertEqual(kwargs["json"]["textQuery"], "Friseursalon")
        self.assertEqual(kwargs["json"]["regionCode"], "DE")
        self.assertIn("locationBias", kwargs["json"])

    def test_discovery_resume_enqueues_children_for_completed_split_tile(self) -> None:
        class FakeClient:
            def __init__(self) -> None:
                self.calls: list[Tile] = []

            def text_search(self, keyword: str, tile: Tile) -> list[dict]:
                self.calls.append(tile)
                return []

        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            checkpoint_path = base / "checkpoints.jsonl"
            raw_dir = base / "raw"
            raw_dir.mkdir()
            root = Tile(52.43, 13.28, 52.57, 13.53)
            checkpoints = CheckpointStore(checkpoint_path)
            checkpoints.write("Friseursalon", root, "done", result_count=60, split=True)
            client = FakeClient()

            discover_google_places(
                ["Friseursalon"],
                [root],
                raw_dir=raw_dir,
                checkpoint_path=checkpoint_path,
                resume=True,
                max_split_depth=1,
                client=client,
            )

            self.assertEqual(len(client.calls), 4)
            self.assertTrue(all(tile.depth == 1 for tile in client.calls))

    def test_smart_split_stops_known_heavy_child_tiles(self) -> None:
        class FakeClient:
            def __init__(self) -> None:
                self.calls: list[Tile] = []

            def text_search(self, keyword: str, tile: Tile) -> list[dict]:
                self.calls.append(tile)
                return [{"id": f"pid-{index}", "displayName": {"text": f"Salon {index}"}} for index in range(60)]

        with tempfile.TemporaryDirectory() as tmp:
            client = FakeClient()
            root = Tile(52.43, 13.28, 52.57, 13.53)
            discover_google_places(
                ["Friseursalon"],
                [root],
                raw_dir=Path(tmp) / "raw",
                checkpoint_path=Path(tmp) / "checkpoints.jsonl",
                resume=True,
                max_split_depth=2,
                dense_threshold=55,
                smart_split=True,
                split_min_new_place_ids=8,
                split_min_new_ratio=0.15,
                client=client,
            )

            self.assertEqual(len(client.calls), 5)
            self.assertEqual([tile.depth for tile in client.calls].count(2), 0)

    def test_keyword_merge_same_place_id(self) -> None:
        records = [rec("pid-1", "Salon A", "Herrenfriseur"), rec("pid-1", "Salon A", "Barbershop")]
        deduped, discarded = deduplicate_records(records, TERMS)
        row = deduped[0].to_row(TERMS)
        self.assertEqual(len(deduped), 1)
        self.assertEqual(len(discarded), 1)
        self.assertEqual(row["erstfund_keyword"], "Herrenfriseur")
        self.assertEqual(row["gefunden_durch_suchbegriffe"], "Herrenfriseur, Barbershop")
        self.assertTrue(row["kw_herrenfriseur"])
        self.assertTrue(row["kw_barbershop"])
        self.assertFalse(row["kw_damenfriseur"])

    def test_keyword_statistics_include_overlap_breakdown(self) -> None:
        raw_records = [
            raw("pid-1", "Salon A", "Herrenfriseur"),
            raw("pid-1", "Salon A", "Barbershop"),
            raw("pid-2", "Salon B", "Herrenfriseur"),
            raw("pid-3", "Salon C", "Herrenfriseur"),
            raw("pid-3", "Salon C", "Barbershop"),
            raw("pid-3", "Salon C", "Damenfriseur"),
        ]
        rows = {row["suchbegriff"]: row for row in build_keyword_statistics(raw_records, CONFIG)}

        herren = rows["Herrenfriseur"]
        self.assertEqual(herren["eindeutige_place_ids"], 3)
        self.assertEqual(herren["exklusive_place_ids"], 1)
        self.assertEqual(herren["auch_ueber_andere_keywords"], 2)
        self.assertEqual(herren["auch_kw_barbershop"], 2)
        self.assertEqual(herren["auch_kw_damenfriseur"], 1)
        self.assertIn("Barbershop: 2", herren["auch_ueber_andere_keywords_details"])
        self.assertIn("Damenfriseur: 1", herren["auch_ueber_andere_keywords_details"])

    def test_load_raw_records_skips_invalid_jsonl_tail(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            raw_dir = Path(tmp)
            with (raw_dir / "broken.jsonl").open("w", encoding="utf-8") as handle:
                handle.write(json.dumps(raw("pid-1", "Salon A", "Herrenfriseur")) + "\n")
                handle.write('{"unfinished": ')

            with self.assertLogs("scraper.pipeline", level="WARNING"):
                self.assertEqual(len(load_raw_records(raw_dir)), 1)

    def test_same_domain_different_addresses_are_not_merged(self) -> None:
        a = rec("pid-a", "Kette Berlin", "Friseursalon", website="https://kette.de")
        b = rec("pid-b", "Kette Hamburg", "Friseursalon", website="https://www.kette.de/kontakt")
        b.strasse, b.hausnummer, b.plz, b.ort = "andere str", "2", "20095", "hamburg"
        deduped, _ = deduplicate_records([a, b], TERMS)
        self.assertEqual(len(deduped), 2)

    def test_place_id_dedup(self) -> None:
        deduped, discarded = deduplicate_records([rec("pid-1", "Salon A"), rec("pid-1", "Salon A")], TERMS)
        self.assertEqual(len(deduped), 1)
        self.assertEqual(discarded[0].aussortiert_grund, "duplikat")

    def test_address_name_strong_match_merges(self) -> None:
        deduped, discarded = deduplicate_records([rec("pid-a", "Marias Friseur"), rec("pid-b", "Maria Friseur")], TERMS)
        self.assertEqual(len(deduped), 1)
        self.assertEqual(discarded[0].aussortiert_grund, "duplikat")

    def test_weak_duplicate_is_manual_not_merged(self) -> None:
        a = rec("pid-a", "Salon Alpha")
        b = rec("pid-b", "Salon Beta")
        b.strasse, b.hausnummer, b.plz, b.ort = "andere str", "2", "20095", "hamburg"
        b.telefon = a.telefon = "030 123456"
        deduped, _ = deduplicate_records([a, b], TERMS)
        self.assertEqual(len(deduped), 2)

    def test_phone_normalization(self) -> None:
        self.assertEqual(normalize_phone("030 123456"), "+4930123456")
        self.assertEqual(normalize_phone("+49 30 123456"), "+4930123456")
        self.assertEqual(normalize_phone("0049 30 123456"), "+4930123456")

    def test_website_normalization(self) -> None:
        self.assertEqual(normalize_url("https://www.beispiel.de/"), "https://beispiel.de")
        self.assertEqual(normalize_domain("http://beispiel.de/kontakt"), "beispiel.de")

    def test_classification(self) -> None:
        clear = classify_record(rec("pid-1", "Friseursalon Klar"))
        unclear = classify_record(PlaceRecord(firmenname="Beauty Studio", google_places_id="pid-2", google_types="beauty_salon"))
        bad = classify_record(PlaceRecord(firmenname="Autohaus", google_places_id="pid-3", google_types="car_repair"))
        self.assertEqual(clear.klassifizierung, "komplett")
        self.assertEqual(unclear.klassifizierung, "manuelle_pruefung")
        self.assertEqual(bad.klassifizierung, "aussortiert")

    def test_xlsx_outputs_from_fixture_pipeline(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            raw_dir = base / "raw" / "google_places"
            raw_dir.mkdir(parents=True)
            records = [
                raw("pid-1", "Friseursalon Eins", "Herrenfriseur", website="https://salon.example"),
                raw("pid-1", "Friseursalon Eins", "Barbershop", website="https://salon.example"),
                raw("pid-2", "Beauty Nails", "Damenfriseur", types=["beauty_salon"]),
                raw("pid-3", "Auto Reparatur", "Friseursalon", types=["car_repair"]),
            ]
            with (raw_dir / "fixture.jsonl").open("w", encoding="utf-8") as handle:
                for item in records:
                    handle.write(json.dumps(item, ensure_ascii=False) + "\n")
            stats = process_raw_to_outputs(
                CONFIG,
                raw_dir=raw_dir,
                processed_dir=base / "processed",
                final_dir=base / "final",
                reports_dir=base / "reports",
                enrich_email=False,
            )
            self.assertEqual(stats["rohtreffer"], 4)
            for filename in ["alles_komplett.xlsx", "manuelle_pruefung.xlsx", "aussortiert.xlsx"]:
                path = base / "final" / filename
                self.assertTrue(path.exists())
                self.assertTrue(verify_workbook(path, ["firmenname", "google_places_id", "kw_herrenfriseur"]))
                wb = load_workbook(path)
                self.assertGreaterEqual(wb.active.max_row, 1)
            self.assertTrue((base / "reports" / "keyword_statistik.xlsx").exists())

    def test_contact_enrichment_checkpoint_reuses_done_website(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            checkpoint = Path(tmp) / "contact.jsonl"
            first = [PlaceRecord(firmenname="Salon", website="https://salon.example")]
            candidate = EmailCandidate("info@salon.example", "https://salon.example/kontakt", 0.9, "allgemein")
            with patch("scraper.enrichment.contacts.find_contact_email", return_value=candidate) as finder:
                enrich_contacts(first, checkpoint_path=checkpoint, timeout=1)
            self.assertEqual(first[0].email, "info@salon.example")
            self.assertEqual(finder.call_count, 1)

            second = [PlaceRecord(firmenname="Salon 2", website="https://salon.example")]
            with patch("scraper.enrichment.contacts.find_contact_email") as finder_again:
                enrich_contacts(second, checkpoint_path=checkpoint, timeout=1)
            self.assertEqual(second[0].email, "info@salon.example")
            finder_again.assert_not_called()

    def test_contact_enrichment_fetches_duplicate_website_once(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            checkpoint = Path(tmp) / "contact.jsonl"
            records = [
                PlaceRecord(firmenname="Salon", website="https://www.salon.example"),
                PlaceRecord(firmenname="Salon Filiale", website="https://salon.example"),
            ]
            candidate = EmailCandidate("info@salon.example", "https://salon.example/kontakt", 0.9, "allgemein")

            with patch("scraper.enrichment.contacts.find_contact_email", return_value=candidate) as finder:
                enrich_contacts(records, checkpoint_path=checkpoint, timeout=1, max_workers=4)

            self.assertEqual(finder.call_count, 1)
            self.assertEqual(records[0].email, "info@salon.example")
            self.assertEqual(records[1].email, "info@salon.example")

    def test_cost_estimate_uses_free_cap(self) -> None:
        estimate = estimate_google_places_cost(text_search_requests=2000)
        self.assertEqual(estimate["text_search_cost_usd"], 35.0)
        self.assertEqual(estimate["details_cost_usd"], 0.0)
        self.assertEqual(estimate["total_cost_usd"], 35.0)


if __name__ == "__main__":
    unittest.main()
